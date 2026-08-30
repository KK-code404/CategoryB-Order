from __future__ import annotations

import io
import json
import secrets
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import settings
from ..models import Dealer, ImportJob, MaterialMapping, SalesOrderLine, Supplier, User
from .audit import add_audit_event

ORDER_HEADERS = ["代理商编码", "订单号", "零件号", "物料号", "品名", "订单数量", "单位", "品牌/供应商编码"]
SHIPMENT_HEADERS = ["订单号", "物料号", "品名", "数量", "收货人", "联系电话", "收货地址", "要求发货日期", "备注"]


# CHANGE [2026-08-30 12:58 +08:00] [WH400]: 在解析前限制大小、文件签名、宏与公式，降低邮件附件带来的风险。
def _open_safe_workbook(content: bytes, filename: str):
    if len(content) > settings.max_attachment_bytes:
        raise ValueError("Excel 文件不能超过 5 MB")
    if not filename.lower().endswith(".xlsx") or not content.startswith(b"PK"):
        raise ValueError("仅支持无宏 .xlsx 文件")
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            if any(cell.data_type == "f" for cell in row):
                raise ValueError("Excel 中不能包含公式")
    return workbook


# CHANGE [2026-08-30 12:58 +08:00] [WH400]: 将表头标准化并阻止缺列模板进入后续业务逻辑。
def _header_map(sheet, expected: list[str]) -> tuple[dict[str, int], int]:
    best_headers: list[str] = []
    for row_no, cells in enumerate(sheet.iter_rows(min_row=1, max_row=10), start=1):
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in cells]
        if all(header in headers for header in expected):
            return {header: headers.index(header) for header in expected}, row_no
        if sum(header in headers for header in expected) > sum(header in best_headers for header in expected):
            best_headers = headers
    missing = [header for header in expected if header not in best_headers]
    raise ValueError(f"缺少必填列：{', '.join(missing)}")


# CHANGE [2026-08-30 12:58 +08:00] [WH400]: 校验订单模板并暂存结构化结果，避免错误行直接污染订单台账。
def preview_order_import(db: Session, actor: User, content: bytes, filename: str) -> ImportJob:
    workbook = _open_safe_workbook(content, filename)
    sheet = workbook.active
    columns, header_row = _header_map(sheet, ORDER_HEADERS)
    valid: list[dict[str, Any]] = []
    errors: list[dict[str, str | int]] = []
    seen: set[tuple[str, str, str]] = set()

    for row_no, cells in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if not any(value is not None for value in cells):
            continue
        raw = {header: cells[index] for header, index in columns.items()}
        try:
            quantity = Decimal(str(raw["订单数量"]))
            if quantity <= 0:
                raise ValueError("订单数量必须大于 0")
            dealer_code = str(raw["代理商编码"] or "").strip()
            order_no = str(raw["订单号"] or "").strip()
            part_no = str(raw["零件号"] or "").strip()
            material_no = str(raw["物料号"] or "").strip()
            supplier_code = str(raw["品牌/供应商编码"] or "").strip()
            if not all([dealer_code, order_no, part_no, material_no, supplier_code]):
                raise ValueError("关键编号不能为空")
            key = (dealer_code, order_no, part_no)
            if key in seen:
                raise ValueError("文件内代理商+订单号+零件号重复")
            seen.add(key)
            if not db.scalar(select(Dealer).where(Dealer.code == dealer_code)):
                raise ValueError("代理商编码不存在")
            if not db.scalar(select(Supplier).where(Supplier.code == supplier_code)):
                raise ValueError("供应商编码不存在")
            valid.append({
                "dealer_code": dealer_code, "order_no": order_no, "part_no": part_no,
                "material_no": material_no, "product_name": str(raw["品名"] or "").strip(),
                "ordered_qty": str(quantity), "unit": str(raw["单位"] or "桶").strip(), "supplier_code": supplier_code,
            })
        except (InvalidOperation, ValueError, TypeError) as exc:
            errors.append({"row": row_no, "message": str(exc)})

    token = secrets.token_hex(24)
    job = ImportJob(token=token, payload_json=json.dumps(valid, ensure_ascii=False), error_json=json.dumps(errors, ensure_ascii=False), created_by=actor.id)
    db.add(job)
    add_audit_event(db, actor=actor, action="订单导入校验", object_type="ImportJob", object_id=token, detail=f"有效 {len(valid)} 行，错误 {len(errors)} 行")
    db.commit()
    return job


# CHANGE [2026-08-30 12:58 +08:00] [WH400]: 只提交无错误的预览任务，并保护已有核销订单数量不被覆盖。
def commit_order_import(db: Session, actor: User, token: str) -> int:
    job = db.scalar(select(ImportJob).where(ImportJob.token == token))
    if not job or job.committed_at:
        raise ValueError("导入任务不存在或已提交")
    if json.loads(job.error_json):
        raise ValueError("导入文件仍有错误，请修正后重新上传")

    count = 0
    for item in json.loads(job.payload_json):
        dealer = db.scalar(select(Dealer).where(Dealer.code == item["dealer_code"]))
        supplier = db.scalar(select(Supplier).where(Supplier.code == item["supplier_code"]))
        assert dealer and supplier
        mapping = db.scalar(select(MaterialMapping).where(MaterialMapping.material_no == item["material_no"]))
        if not mapping:
            db.add(MaterialMapping(material_no=item["material_no"], part_no=item["part_no"], product_name=item["product_name"], brand_code=item["supplier_code"], supplier_id=supplier.id))
        existing = db.scalar(select(SalesOrderLine).where(SalesOrderLine.dealer_id == dealer.id, SalesOrderLine.order_no == item["order_no"], SalesOrderLine.part_no == item["part_no"]))
        quantity = Decimal(item["ordered_qty"])
        if existing:
            if existing.reconciled_qty and quantity != existing.ordered_qty:
                raise ValueError(f"订单 {item['order_no']} 已发生核销，禁止修改订单数量")
            existing.ordered_qty = quantity
            existing.product_name = item["product_name"]
            existing.unit = item["unit"]
            existing.version += 1
        else:
            db.add(SalesOrderLine(dealer_id=dealer.id, order_no=item["order_no"], part_no=item["part_no"], material_no=item["material_no"], product_name=item["product_name"], ordered_qty=quantity, reconciled_qty=0, unit=item["unit"]))
        count += 1

    job.committed_at = datetime.now(timezone.utc)
    add_audit_event(db, actor=actor, action="订单导入提交", object_type="ImportJob", object_id=token, detail=f"写入订单台账 {count} 行")
    db.commit()
    return count


# CHANGE [2026-08-30 12:58 +08:00] [WH400]: 解析标准发货工作表并保留原始行号，供邮件入库与错误定位使用。
def parse_shipment_workbook(content: bytes, filename: str) -> list[dict[str, Any]]:
    workbook = _open_safe_workbook(content, filename)
    if "发货申请" not in workbook.sheetnames:
        raise ValueError("工作表名称必须为“发货申请”")
    sheet = workbook["发货申请"]
    columns, header_row = _header_map(sheet, SHIPMENT_HEADERS)
    rows: list[dict[str, Any]] = []
    for row_no, cells in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if not any(value is not None for value in cells):
            continue
        item = {header: cells[index] for header, index in columns.items()}
        item["_row"] = row_no
        rows.append(item)
    return rows
